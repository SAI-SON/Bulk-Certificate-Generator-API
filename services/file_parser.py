import csv
import io
import re
from datetime import datetime
from openpyxl import load_workbook
from django.conf import settings
from django.core.validators import validate_email
from django.core.exceptions import ValidationError


REQUIRED_COLUMNS = ["name"]
OPTIONAL_COLUMNS = ["email", "course_name", "completion_date"]


class FileParserError(Exception):
    pass


class ParticipantFileParser:
    """
    Parses and validates CSV and XLSX files containing participant data.
    """

    def __init__(self, file_path):
        self.file_path = file_path

    def parse(self):
        """
        Detect file type and parse accordingly.
        Returns a tuple: (list of dicts representing valid rows, list of error dicts, integer total_rows)
        """
        ext = str(self.file_path).split(".")[-1].lower()
        if ext == "csv":
            return self._parse_csv()
        elif ext == "xlsx":
            return self._parse_xlsx()
        else:
            raise FileParserError(f"Unsupported file format: {ext}")

    def _parse_csv(self):
        valid_rows = []
        errors = []
        
        try:
            # We open with utf-8 encoding
            with open(self.file_path, mode="r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                
                # Validate headers
                if not reader.fieldnames:
                    raise FileParserError("The CSV file is empty or has no headers.")
                
                self._validate_headers(reader.fieldnames)
                
                row_count = 0
                for i, row in enumerate(reader):
                    row_count += 1
                    if row_count > settings.MAX_PARTICIPANT_ROWS:
                        raise FileParserError(f"Row count exceeds the maximum allowed ({settings.MAX_PARTICIPANT_ROWS}).")
                    
                    cleaned_row = self._clean_row(row)
                    row_errors = self._validate_row(cleaned_row, i + 2)  # +2 for 1-based index and header
                    
                    if row_errors:
                        errors.extend(row_errors)
                    else:
                        valid_rows.append(cleaned_row)
                        
            # Duplicate detection based on name and email
            self._detect_duplicates(valid_rows, errors)
            
            return valid_rows, errors, row_count
            
        except FileParserError:
            raise
        except UnicodeDecodeError:
            raise FileParserError("The CSV file must be UTF-8 encoded.")
        except Exception as e:
            raise FileParserError(f"Error reading CSV file: {str(e)}")

    def _parse_xlsx(self):
        valid_rows = []
        errors = []
        
        try:
            wb = load_workbook(filename=self.file_path, read_only=True, data_only=True)
            sheet = wb.active
            
            headers = []
            row_count = 0
            
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i == 0:
                    headers = [str(cell).strip() if cell else "" for cell in row]
                    self._validate_headers(headers)
                    continue
                    
                row_count += 1
                if row_count > settings.MAX_PARTICIPANT_ROWS:
                    raise FileParserError(f"Row count exceeds the maximum allowed ({settings.MAX_PARTICIPANT_ROWS}).")
                
                # Check if row is entirely empty
                if not any(row):
                    continue
                    
                row_dict = dict(zip(headers, row))
                cleaned_row = self._clean_row(row_dict)
                row_errors = self._validate_row(cleaned_row, i + 1)
                
                if row_errors:
                    errors.extend(row_errors)
                else:
                    valid_rows.append(cleaned_row)
                    
            wb.close()
            
            self._detect_duplicates(valid_rows, errors)
            
            return valid_rows, errors, row_count
            
        except FileParserError:
            raise
        except Exception as e:
            raise FileParserError(f"Error reading XLSX file: {str(e)}")

    def _validate_headers(self, headers):
        headers_lower = [h.lower() for h in headers if h]
        for req in REQUIRED_COLUMNS:
            if req not in headers_lower:
                raise FileParserError(f"Missing required column: '{req}'")

    def _clean_row(self, row):
        """Trim whitespace and convert keys to lowercase."""
        cleaned = {}
        for k, v in row.items():
            if not k:
                continue
            
            key_lower = k.lower()
            if key_lower in REQUIRED_COLUMNS + OPTIONAL_COLUMNS:
                if isinstance(v, str):
                    cleaned[key_lower] = v.strip()
                elif v is not None:
                    # Handle numbers or dates in XLSX
                    if isinstance(v, datetime):
                        cleaned[key_lower] = v.strftime("%Y-%m-%d")
                    else:
                        cleaned[key_lower] = str(v).strip()
                else:
                    cleaned[key_lower] = ""
        return cleaned

    def _validate_row(self, row, line_number):
        errors = []
        
        # Name is required
        if not row.get("name"):
            errors.append({"line": line_number, "error": "Name is required."})
            
        # Email validation (optional)
        email = row.get("email")
        if email:
            try:
                validate_email(email)
            except ValidationError:
                errors.append({"line": line_number, "error": f"Invalid email address: {email}"})
                
        # Date validation (optional)
        date_str = row.get("completion_date")
        if date_str:
            try:
                # Try standard formats
                found_valid = False
                for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
                    try:
                        parsed = datetime.strptime(date_str, fmt)
                        # Normalize to YYYY-MM-DD
                        row["completion_date"] = parsed.strftime("%Y-%m-%d")
                        found_valid = True
                        break
                    except ValueError:
                        pass
                
                if not found_valid:
                     errors.append({"line": line_number, "error": f"Invalid date format (use YYYY-MM-DD): {date_str}"})
            except Exception:
                errors.append({"line": line_number, "error": f"Invalid date: {date_str}"})
                
        return errors

    def _detect_duplicates(self, valid_rows, errors):
        """Detect duplicates based on name and email combination to avoid generating identical certificates."""
        seen = set()
        duplicates_to_remove = []
        
        for i, row in enumerate(valid_rows):
            name = row.get("name", "").lower()
            email = row.get("email", "").lower()
            
            identifier = f"{name}:{email}"
            
            if identifier in seen:
                # Find original index for error reporting
                # (approximate since we don't carry line numbers in valid_rows currently)
                errors.append({"line": "multiple", "error": f"Duplicate record found for {row.get('name')} ({email})"})
                duplicates_to_remove.append(i)
            else:
                seen.add(identifier)
                
        # Remove duplicates from valid_rows (in reverse order to not mess up indices)
        for i in reversed(duplicates_to_remove):
            valid_rows.pop(i)
